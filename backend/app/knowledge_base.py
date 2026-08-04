from __future__ import annotations

import csv
import hashlib
import io
import json
import re
import threading
import tempfile
import uuid
from datetime import date, datetime, timezone
from decimal import Decimal
from pathlib import Path
from typing import Any

import xlrd
from docx import Document as WordDocument
from langchain_chroma import Chroma
from langchain_core.documents import Document
from langchain_openai import ChatOpenAI, OpenAIEmbeddings
from langchain_text_splitters import RecursiveCharacterTextSplitter
from openpyxl import load_workbook
from pypdf import PdfReader

from .config import ROOT_DIR, settings
from .database import (
    complete_knowledge_document,
    count_knowledge_records,
    create_knowledge_document,
    delete_knowledge_document_record,
    fail_knowledge_document,
    get_knowledge_document_by_hash,
    get_knowledge_records_by_ids,
    count_knowledge_records_for_user,
    list_knowledge_documents,
    list_knowledge_records,
    replace_knowledge_records,
    search_knowledge_structured_records,
)

ALLOWED_EXTENSIONS = {".pdf", ".xls", ".xlsx", ".csv", ".txt", ".md", ".docx"}
MIME_TYPES = {
    ".pdf": {"application/pdf"},
    ".docx": {"application/vnd.openxmlformats-officedocument.wordprocessingml.document", "application/zip"},
    ".xlsx": {"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "application/zip"},
    ".xls": {"application/vnd.ms-excel", "application/octet-stream"},
    ".csv": {"text/csv", "text/plain", "application/csv", "application/vnd.ms-excel"},
    ".txt": {"text/plain"},
    ".md": {"text/markdown", "text/plain", "text/x-markdown"},
}
_store: Chroma | None = None
_store_lock = threading.Lock()


def _company_id(user_id: int) -> str:
    """The current schema is single-user tenancy, so the authenticated user is the company boundary."""
    return str(user_id)


def _chroma_path() -> Path:
    requested = Path(settings.chroma_persist_directory)
    try:
        requested.mkdir(parents=True, exist_ok=True)
        probe = requested / ".write-test"
        probe.touch(exist_ok=True)
        probe.unlink(missing_ok=True)
        return requested
    except PermissionError:
        # Local development fallback; production should grant access to the configured EC2 directory.
        fallback = Path(tempfile.gettempdir()) / "company-rag" / "chroma"
        fallback.mkdir(parents=True, exist_ok=True)
        return fallback


def _vector_store() -> Chroma:
    global _store
    if not settings.openai_api_key:
        raise ValueError("OPENAI_API_KEY is not configured.")
    if _store is None:
        with _store_lock:
            if _store is None:
                _store = Chroma(
                    collection_name=settings.chroma_collection_name,
                    embedding_function=OpenAIEmbeddings(model="text-embedding-3-small", api_key=settings.openai_api_key),
                    persist_directory=str(_chroma_path()),
                    collection_metadata={"hnsw:space": "cosine"},
                )
    return _store


def _clean(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, float) and value != value:
        return ""
    if isinstance(value, Decimal):
        value = format(value, "f")
    return re.sub(r"\s+", " ", str(value)).strip()


def _headers(values: list[Any]) -> list[str]:
    result: list[str] = []
    counts: dict[str, int] = {}
    for index, value in enumerate(values, start=1):
        base = _clean(value) or f"Column {index}"
        base = re.sub(r"[\x00-\x1f]+", " ", base).strip()[:120]
        counts[base] = counts.get(base, 0) + 1
        result.append(base if counts[base] == 1 else f"{base}_{counts[base]}")
    return result


def _header_index(rows: list[list[Any]]) -> int:
    candidates = rows[: min(20, len(rows))]
    scores: list[float] = []
    for row in candidates:
        nonempty = [_clean(value) for value in row if _clean(value)]
        strings = sum(not value.replace(".", "", 1).isdigit() for value in nonempty)
        uniqueness = len(set(nonempty)) / max(1, len(nonempty))
        # Headers are normally mostly text, unique, and followed by similarly
        # populated rows. A one-cell report title receives a strong penalty.
        next_widths = [sum(bool(_clean(value)) for value in following) for following in candidates[len(scores) + 1:len(scores) + 4]]
        continuity = sum(abs(len(nonempty) - width) <= 1 for width in next_widths)
        scores.append(len(nonempty) * 2 + strings * 1.5 + uniqueness + continuity * 2 - (8 if len(nonempty) == 1 else 0))
    return max(range(len(scores)), key=scores.__getitem__) if scores else 0


def _record(content: str, structured: dict[str, str], metadata: dict[str, Any]) -> dict[str, Any]:
    safe_metadata = {key: (value if isinstance(value, (str, int, float, bool)) and value is not None else "") for key, value in metadata.items()}
    return {"content": content.strip(), "structured_data": structured, "metadata": safe_metadata}


def _table_records(rows: list[list[Any]], filename: str, file_type: str, sheet_name: str = "") -> list[dict[str, Any]]:
    # Keep physical row numbers while removing empty rows from consideration.
    indexed_rows = [(number, list(row)) for number, row in enumerate(rows, start=1) if any(_clean(value) for value in row)]
    if not indexed_rows:
        return []
    raw_rows = [row for _, row in indexed_rows]
    used_columns = [index for index in range(max(map(len, raw_rows))) if any(index < len(row) and _clean(row[index]) for row in raw_rows)]
    cleaned_rows = [[row[index] if index < len(row) else None for index in used_columns] for row in raw_rows]
    header_row = _header_index(cleaned_rows)
    headers = _headers(cleaned_rows[header_row])
    result = []
    normalized_header = [_clean(value).casefold() for value in cleaned_rows[header_row]]
    is_serial_table = bool(normalized_header and any(token in normalized_header[0] for token in ("sl", "serial", "sr.")))
    grouped: list[tuple[int, list[dict[str, str]]]] = []
    current_group: tuple[int, list[dict[str, str]]] | None = None
    for offset, row in enumerate(cleaned_rows[header_row + 1 :], start=header_row + 1):
        data_index = indexed_rows[offset][0]
        # Ignore repeated headers commonly inserted between printed page blocks.
        if [_clean(value).casefold() for value in row] == normalized_header:
            continue
        structured = {headers[index]: _clean(value) for index, value in enumerate(row) if index < len(headers) and _clean(value)}
        first_value = _clean(row[0]) if row else ""
        if is_serial_table and structured:
            if re.fullmatch(r"\d+(?:\.0+)?", first_value):
                current_group = (data_index, [structured])
                grouped.append(current_group)
            elif current_group:
                current_group[1].append(structured)
            continue
        if structured:
            result.append(_record("\n".join(f"{key}: {value}" for key, value in structured.items()), structured, {
                "source": filename, "file_type": file_type, "sheet_name": sheet_name,
                "page_number": 0, "row_number": data_index, "chunk_index": 0,
            }))
    for row_number, group_rows in grouped:
        grouped_data = {"rows": group_rows}
        readable = "\n".join(
            " | ".join(f"{key}: {value}" for key, value in row.items()) for row in group_rows
        )
        result.append(_record(readable, grouped_data, {
            "source": filename, "file_type": file_type, "sheet_name": sheet_name,
            "page_number": 0, "row_number": row_number, "chunk_index": 0,
        }))
    return result


def _excel_records(path: Path, filename: str) -> list[dict[str, Any]]:
    records: list[dict[str, Any]] = []
    workbook = load_workbook(path, read_only=True, data_only=True)
    for sheet in workbook.worksheets:
        try:
            records.extend(_table_records([list(row) for row in sheet.iter_rows(values_only=True)], filename, "xlsx", sheet.title))
        except Exception:
            continue
    return records


def _xls_records(path: Path, filename: str) -> list[dict[str, Any]]:
    records: list[dict[str, Any]] = []
    workbook = xlrd.open_workbook(str(path))
    for sheet in workbook.sheets():
        try:
            rows = [[sheet.cell_value(row, column) for column in range(sheet.ncols)] for row in range(sheet.nrows)]
            records.extend(_table_records(rows, filename, "xls", sheet.name))
        except Exception:
            continue
    return records


def _decode(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-16", "cp1252", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _csv_records(path: Path, filename: str) -> list[dict[str, Any]]:
    text = _decode(path.read_bytes())
    sample = text[:8192]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    return _table_records(list(csv.reader(io.StringIO(text), dialect)), filename, "csv")


def _split_records(text: str, filename: str, file_type: str, **location: Any) -> list[dict[str, Any]]:
    splitter = RecursiveCharacterTextSplitter(chunk_size=1200, chunk_overlap=180, separators=["\n\n", "\n", ". ", " "])
    return [_record(chunk, {}, {"source": filename, "file_type": file_type, "sheet_name": "", "page_number": 0,
             "row_number": 0, "chunk_index": index, **location})
            for index, chunk in enumerate(splitter.split_text(text.strip())) if chunk.strip()]


def extract_records(path: Path, filename: str) -> list[dict[str, Any]]:
    extension = path.suffix.lower()
    if extension == ".xlsx":
        return _excel_records(path, filename)
    if extension == ".xls":
        return _xls_records(path, filename)
    if extension == ".csv":
        return _csv_records(path, filename)
    if extension == ".pdf":
        records = []
        for page_number, page in enumerate(PdfReader(str(path)).pages, start=1):
            text = (page.extract_text() or "").strip()
            if text:
                records.extend(_split_records(text, filename, "pdf", page_number=page_number))
        return records
    if extension == ".docx":
        word = WordDocument(str(path))
        sections: list[str] = []
        current: list[str] = []
        for paragraph in word.paragraphs:
            text = paragraph.text.strip()
            if not text:
                continue
            if paragraph.style and paragraph.style.name.lower().startswith("heading") and current:
                sections.append("\n".join(current)); current = []
            current.append(text)
        if current:
            sections.append("\n".join(current))
        return [record for section_index, section in enumerate(sections) for record in _split_records(section, filename, "docx", section_index=section_index)]
    return _split_records(_decode(path.read_bytes()), filename, extension.lstrip("."))


def validate_upload(filename: str, content_type: str) -> tuple[str, str]:
    safe_name = re.sub(r"[^A-Za-z0-9._() -]+", "_", Path(filename or "document").name).strip(". ") or "document"
    extension = Path(safe_name).suffix.lower()
    if extension not in ALLOWED_EXTENSIONS:
        raise ValueError(f"Unsupported file type: {extension or 'unknown'}")
    normalized_type = content_type.split(";", 1)[0].lower().strip()
    if normalized_type and normalized_type not in MIME_TYPES[extension] and normalized_type != "application/octet-stream":
        raise ValueError(f"The MIME type {normalized_type} does not match {extension}.")
    return safe_name, extension


def _vector_id(company_id: str, file_id: str, metadata: dict[str, Any]) -> str:
    location = metadata.get("sheet_name") or (f"page_{metadata.get('page_number')}" if metadata.get("page_number") else None) or (f"section_{metadata.get('section_index')}" if metadata.get("section_index") not in (None, "") else "text")
    position = metadata.get("row_number") or f"chunk_{metadata.get('chunk_index', 0)}"
    return f"{company_id}:{file_id}:{location}:{position}"


def ingest_file(user_id: int, filename: str, path: Path, content_type: str = "") -> dict[str, Any]:
    safe_name, extension = validate_upload(filename, content_type)
    file_size = path.stat().st_size
    if not file_size:
        raise ValueError("The uploaded file is empty.")
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    existing = get_knowledge_document_by_hash(user_id, digest)
    if existing and existing.get("status") in {"ready", "completed"} and count_knowledge_records(existing["id"], user_id):
        return {"id": existing["id"], "filename": existing["filename"], "chunks": existing["chunk_count"],
                "uploaded_at": existing["created_at"].isoformat(), "duplicate": True}
    if existing:
        delete_document(user_id, existing["id"])
    records = extract_records(path, safe_name)
    if not records:
        raise ValueError("No readable text was found. Scanned PDFs require OCR, which is not available.")
    file_id = uuid.uuid4().hex
    company_id = _company_id(user_id)
    created_at = datetime.now(timezone.utc).isoformat()
    create_knowledge_document(file_id, user_id, safe_name, content_type, digest, file_size, extension.lstrip("."))
    documents: list[Document] = []
    ids: list[str] = []
    for index, record in enumerate(records):
        metadata = record["metadata"]
        metadata.update({"company_id": company_id, "user_id": str(user_id), "file_id": file_id})
        record_id = _vector_id(company_id, file_id, metadata)
        record["record_id"] = record_id
        chroma_metadata = {**metadata, "record_id": record_id}
        documents.append(Document(page_content=record["content"], metadata=chroma_metadata))
        ids.append(record_id)
    replace_knowledge_records(file_id, user_id, records)
    try:
        store = _vector_store()
        batch_size = max(1, settings.knowledge_embedding_batch_size)
        for start in range(0, len(documents), batch_size):
            store.add_documents(documents[start:start + batch_size], ids=ids[start:start + batch_size])
        complete_knowledge_document(file_id, user_id, len(records), len(documents))
        return {"id": file_id, "filename": safe_name, "chunks": len(documents), "records": len(records),
                "uploaded_at": created_at, "duplicate": False}
    except Exception as exc:
        try:
            _vector_store().delete(where={"$and": [{"company_id": company_id}, {"file_id": file_id}]})
        except Exception:
            pass
        fail_knowledge_document(file_id, user_id, str(exc))
        raise


def list_documents(user_id: int) -> list[dict[str, Any]]:
    documents = list_knowledge_documents(user_id)
    for document in documents:
        if hasattr(document.get("uploaded_at"), "isoformat"):
            document["uploaded_at"] = document["uploaded_at"].isoformat()
    return documents


def get_records(user_id: int, document_ids: list[str] | None = None, limit: int | None = None) -> list[dict[str, Any]]:
    return list_knowledge_records(user_id, document_ids, limit)


def delete_document(user_id: int, document_id: str) -> bool:
    company_id = _company_id(user_id)
    owned = any(item["id"] == document_id for item in list_knowledge_documents(user_id))
    if not owned:
        return False
    _vector_store().delete(where={"$and": [{"company_id": company_id}, {"file_id": document_id}]})
    return delete_knowledge_document_record(document_id, user_id)


def _source_for_record(record: dict[str, Any]) -> dict[str, Any]:
    metadata = record.get("metadata") or {}
    return {
        "file_id": metadata.get("file_id") or record.get("document_id"),
        "record_id": record.get("record_id"),
        "source": metadata.get("source"),
        "sheet_name": metadata.get("sheet_name") or None,
        "page_number": metadata.get("page_number") or None,
        "row_number": metadata.get("row_number") or None,
    }


def _original_record_data(record: dict[str, Any]) -> dict[str, Any]:
    """Return lossless source JSON, including compatibility with older uploads."""
    data = record.get("structured_data") or {}
    if isinstance(data, dict) and isinstance(data.get("source_fields"), dict):
        return data["source_fields"]
    return data if isinstance(data, dict) else {}


def _query_terms(question: str) -> tuple[list[str], list[str]]:
    stop_words = {
        "a", "all", "am", "and", "are", "company", "details", "for", "from",
        "give", "in", "is", "me", "of", "on", "show", "the", "to", "with",
    }
    terms = [
        token for token in re.findall(r"[a-z0-9@._+-]+", question.casefold())
        if len(token) > 1 and token not in stop_words
    ]
    identifiers = [term for term in terms if (term.isdigit() and len(term) >= 4) or "@" in term]
    return terms, identifiers


def _retrieve_records(
    user_id: int,
    question: str,
    document_ids: list[str] | None,
) -> list[dict[str, Any]]:
    """Retrieve from the full tenant index without loading every row per query."""
    _, identifier_terms = _query_terms(question)
    if identifier_terms:
        candidates = search_knowledge_structured_records(
            user_id, identifier_terms, document_ids, limit=5000
        )
        records = []
        for record in candidates:
            searchable = json.dumps(_original_record_data(record), ensure_ascii=False).casefold()
            if all(
                re.search(rf"(?<![a-z0-9]){re.escape(term)}(?![a-z0-9])", searchable)
                for term in identifier_terms
            ):
                records.append(record)
        return _deduplicate_records(records, limit=250)

    company_id = _company_id(user_id)
    where: dict[str, Any] = {"company_id": company_id}
    if document_ids:
        where = {
            "$and": [
                {"company_id": {"$eq": company_id}},
                {"file_id": {"$in": document_ids}},
            ]
        }
    retrieved = _vector_store().similarity_search(question, k=60, filter=where)
    record_ids = [str(document.metadata.get("record_id") or "") for document in retrieved]
    by_id = get_knowledge_records_by_ids(user_id, [record_id for record_id in record_ids if record_id])
    records = [by_id[record_id] for record_id in record_ids if record_id in by_id]
    return _deduplicate_records(records, limit=60)


def _deduplicate_records(records: list[dict[str, Any]], limit: int) -> list[dict[str, Any]]:
    unique: list[dict[str, Any]] = []
    seen_data: set[str] = set()
    for record in records:
        fingerprint = json.dumps(_original_record_data(record), ensure_ascii=False, sort_keys=True)
        if fingerprint not in seen_data:
            unique.append(record)
            seen_data.add(fingerprint)
        if len(unique) == limit:
            break
    return unique


def answer_question(
    user_id: int,
    question: str,
    document_ids: list[str] | None = None,
    history: list[dict[str, Any]] | None = None,
) -> dict[str, Any]:
    question = question.strip()
    if not question:
        raise ValueError("Question is required.")
    lower_question = question.casefold()
    _, identifiers = _query_terms(question)
    if not identifiers and re.search(r"\b(?:how many|count of|number of)\s+(?:stored\s+)?(?:records|rows|entries)\b", lower_question):
        total = count_knowledge_records_for_user(user_id, document_ids)
        return {
            "answer": f"| Metric | Result |\n|---|---:|\n| Stored records | {total} |",
            "sources": [],
        }
    records = _retrieve_records(user_id, question, document_ids)
    if not records:
        return {
            "answer": "| Status | Details |\n|---|---|\n| No match | No matching record was found in the stored documents. |",
            "sources": [],
        }
    record_contexts = [
        {
            "record_id": record["record_id"],
            "source": record.get("metadata", {}).get("source"),
            "sheet_name": record.get("metadata", {}).get("sheet_name"),
            "row_number": record.get("metadata", {}).get("row_number"),
            "data": _original_record_data(record),
        }
        for record in records
    ]
    sources = [_source_for_record(record) for record in records]
    recent_history = [
        {"role": str(item.get("role") or ""), "content": str(item.get("content") or "")[:3000]}
        for item in (history or [])[-12:]
        if item.get("role") in {"user", "assistant"} and item.get("content")
    ]
    model = ChatOpenAI(model=settings.openai_model, api_key=settings.openai_api_key, temperature=0)
    response = model.invoke([
        ("system", "You are a retrieval-augmented knowledge-base assistant. Answer only from the supplied original document records. Different files can have different JSON fields; use the fields exactly as supplied and never assume missing facts. Always return exactly one valid Markdown table, regardless of how the user phrases the request. Put one matching record per row and use concise, relevant columns. For a count or summary, use Metric and Result columns. If no evidence answers the question, use Status and Details columns and clearly state that no matching record was found. Escape any literal pipe inside a cell as \\|. Do not add prose before the table. Mention source filename, sheet, or row in a Source column when useful. Keep the result direct and responsive."),
        ("human", f"Conversation history:\n{json.dumps(recent_history, ensure_ascii=False)}\n\nQuestion: {question}\n\nRetrieved original records:\n{json.dumps(record_contexts, ensure_ascii=False)}"),
    ])
    return {"answer": str(response.content), "sources": sources}


def knowledge_status(user_id: int) -> dict[str, Any]:
    documents = list_documents(user_id)
    return {"ready": bool(settings.openai_api_key), "embedding_model": "text-embedding-3-small",
            "answer_model": settings.openai_model, "retrieval": "langchain-chroma",
            "documents": len(documents),
            "chunks": sum(item["chunks"] for item in documents), "supported_extensions": sorted(ALLOWED_EXTENSIONS),
            "max_upload_mb": settings.knowledge_max_upload_mb}
